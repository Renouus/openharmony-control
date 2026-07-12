# -*- coding: utf-8 -*-
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSETS_DIR = DOCS_DIR / "generated_design_assets"
OUTPUT_PATH = DOCS_DIR / "智能家居控制系统产品总体设计文档.docx"

BLACK = (32, 32, 32)
GRAY = (92, 92, 92)
LIGHT_GRAY = (244, 246, 248)
LINE_GRAY = (165, 165, 165)
ACCENT = (64, 108, 176)
ACCENT_LIGHT = (228, 236, 247)
GREEN = (84, 130, 53)
GREEN_LIGHT = (229, 239, 218)
ORANGE = (191, 96, 36)
ORANGE_LIGHT = (248, 228, 214)


@dataclass
class FigureInfo:
    key: str
    title: str
    path: Path


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def load_font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines() or [""]:
        current = ""
        for ch in raw_line:
            candidate = current + ch
            if draw.textlength(candidate, font=font) <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = ch
        if current:
            lines.append(current)
        elif raw_line == "":
            lines.append("")
    return lines


def add_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    title_font: ImageFont.FreeTypeFont,
    body_font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int] = (250, 250, 250),
    outline: tuple[int, int, int] = BLACK,
    radius: int = 18,
) -> None:
    draw.rounded_rectangle(xy, radius=radius, outline=outline, width=3, fill=fill)
    left, top, right, bottom = xy
    draw.text((left + 16, top + 12), title, fill=BLACK, font=title_font)
    text_y = top + 52
    for line in wrap_text(draw, body, body_font, right - left - 32):
        draw.text((left + 16, text_y), line, fill=GRAY, font=body_font)
        text_y += body_font.size + 8


def add_small_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    *,
    font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int] = ACCENT_LIGHT,
) -> None:
    draw.rounded_rectangle(xy, radius=14, outline=BLACK, width=2, fill=fill)
    left, top, right, bottom = xy
    lines = wrap_text(draw, text, font, right - left - 20)
    total_h = len(lines) * font.size + max(len(lines) - 1, 0) * 6
    y = top + ((bottom - top) - total_h) // 2
    for line in lines:
        width = int(draw.textlength(line, font=font))
        x = left + ((right - left) - width) // 2
        draw.text((x, y), line, fill=BLACK, font=font)
        y += font.size + 6


def add_database_symbol(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    title_font: ImageFont.FreeTypeFont,
    body_font: ImageFont.FreeTypeFont,
) -> None:
    left, top, right, bottom = xy
    draw.rectangle((left, top + 20, right, bottom - 20), outline=BLACK, width=3, fill=LIGHT_GRAY)
    draw.ellipse((left, top, right, top + 40), outline=BLACK, width=3, fill=(255, 255, 255))
    draw.ellipse((left, bottom - 40, right, bottom), outline=BLACK, width=3, fill=LIGHT_GRAY)
    draw.text((left + 18, top + 48), title, fill=BLACK, font=title_font)
    y = top + 92
    for line in wrap_text(draw, body, body_font, right - left - 32):
        draw.text((left + 18, y), line, fill=GRAY, font=body_font)
        y += body_font.size + 8


def draw_arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    label: str | None = None,
    font: ImageFont.FreeTypeFont | None = None,
    color: tuple[int, int, int] = BLACK,
    width: int = 3,
) -> None:
    draw.line([start, end], fill=color, width=width)
    dx = end[0] - start[0]
    dy = end[1] - start[1]
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux = dx / length
    uy = dy / length
    arrow = 14
    left = (int(end[0] - arrow * ux - 6 * uy), int(end[1] - arrow * uy + 6 * ux))
    right = (int(end[0] - arrow * ux + 6 * uy), int(end[1] - arrow * uy - 6 * ux))
    draw.polygon([end, left, right], fill=color)
    if label and font:
        mid = ((start[0] + end[0]) // 2, (start[1] + end[1]) // 2)
        bbox = draw.textbbox((0, 0), label, font=font)
        rect = (
            mid[0] - (bbox[2] - bbox[0]) // 2 - 10,
            mid[1] - (bbox[3] - bbox[1]) // 2 - 8,
            mid[0] + (bbox[2] - bbox[0]) // 2 + 10,
            mid[1] + (bbox[3] - bbox[1]) // 2 + 8,
        )
        draw.rounded_rectangle(rect, radius=10, outline=LINE_GRAY, fill=(255, 255, 255))
        draw.text((rect[0] + 10, rect[1] + 6), label, fill=BLACK, font=font)


def draw_dashed_line(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: tuple[int, int, int] = LINE_GRAY,
    width: int = 2,
    step: int = 12,
) -> None:
    dx = end[0] - start[0]
    dy = end[1] - start[1]
    length = int((dx * dx + dy * dy) ** 0.5)
    if length == 0:
        return
    for i in range(0, length, step * 2):
        s_ratio = i / length
        e_ratio = min(i + step, length) / length
        sx = int(start[0] + dx * s_ratio)
        sy = int(start[1] + dy * s_ratio)
        ex = int(start[0] + dx * e_ratio)
        ey = int(start[1] + dy * e_ratio)
        draw.line((sx, sy, ex, ey), fill=color, width=width)


def save_context_diagram(path: Path) -> None:
    img = Image.new("RGB", (1680, 980), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    body_font = load_font(22)
    label_font = load_font(20)

    draw.text((64, 42), "系统上下文图", fill=BLACK, font=title_font)

    add_box(
        draw,
        (610, 280, 1080, 600),
        "OmniHome 智能家居控制系统",
        "统一提供设备状态展示、控制命令下发、场景执行、自动化管理、历史追踪与实时同步能力。",
        title_font=title_font,
        body_font=body_font,
        fill=ACCENT_LIGHT,
    )
    add_box(draw, (80, 120, 400, 300), "家庭用户", "查看家庭总览、控制灯光空调、执行场景、接收通知。", title_font=title_font, body_font=body_font)
    add_box(draw, (80, 520, 400, 700), "家庭成员", "关注日常使用入口，如门禁、广播、环境调节与消息提醒。", title_font=title_font, body_font=body_font)
    add_box(draw, (1240, 120, 1560, 300), "OpenHarmony 客户端", "ArkTS 页面、ViewModel、Repository 与本地缓存。", title_font=title_font, body_font=body_font, fill=GREEN_LIGHT)
    add_box(draw, (1240, 360, 1560, 540), "控制中心服务", "Fastify 路由、领域服务、签名校验、同步与 WebSocket 广播。", title_font=title_font, body_font=body_font, fill=GREEN_LIGHT)
    add_box(draw, (1240, 600, 1560, 780), "设备与数据基础设施", "SQLite、模拟设备、后续可替换的真实 IoT 适配层。", title_font=title_font, body_font=body_font, fill=GREEN_LIGHT)

    draw_arrow(draw, (400, 210), (610, 360), label="使用与反馈", font=label_font)
    draw_arrow(draw, (400, 610), (610, 520), label="家庭协同", font=label_font)
    draw_arrow(draw, (1080, 360), (1240, 210), label="交互入口", font=label_font)
    draw_arrow(draw, (1080, 440), (1240, 450), label="业务服务", font=label_font)
    draw_arrow(draw, (1080, 520), (1240, 690), label="状态来源", font=label_font)

    img.save(path)


def save_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1720, 1040), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    box_title_font = load_font(26, bold=True)
    body_font = load_font(21)
    label_font = load_font(19)

    draw.text((60, 36), "系统总体架构图", fill=BLACK, font=title_font)
    draw.rounded_rectangle((40, 90, 1660, 980), radius=26, outline=BLACK, width=3)

    add_box(draw, (140, 150, 1550, 270), "表现层 / OpenHarmony 客户端", "首页总览、门禁、摄像头、气候、家庭、场景、自动化与通知页面；负责展示、交互与本地状态呈现。", title_font=box_title_font, body_font=body_font, fill=ACCENT_LIGHT)
    add_box(draw, (140, 330, 1550, 480), "应用服务层 / 控制中心服务", "Fastify 路由、场景管理、自动化规则、命令执行、签名校验、历史记录、同步接口与 WebSocket 事件。", title_font=box_title_font, body_font=body_font, fill=GREEN_LIGHT)
    add_box(draw, (140, 540, 760, 790), "共享契约层", "统一设备模型、命令结构、场景定义、安全信封与响应结构，保证前后端语义一致。", title_font=box_title_font, body_font=body_font, fill=ORANGE_LIGHT)
    add_database_symbol(draw, (920, 540, 1200, 820), "SQLite 数据库", "保存 metadata、devices、rooms、scenes、automations 与 history。", title_font=box_title_font, body_font=body_font)
    add_box(draw, (1280, 540, 1550, 820), "模拟设备 / 适配层", "门锁、灯光、空调、环境数据与后续真实 IoT 网关替换点。", title_font=box_title_font, body_font=body_font)

    draw_arrow(draw, (850, 270), (850, 330), label="HTTP / WebSocket", font=label_font)
    draw_arrow(draw, (500, 480), (500, 540), label="类型约束", font=label_font)
    draw_arrow(draw, (1060, 480), (1060, 540), label="持久化", font=label_font)
    draw_arrow(draw, (1410, 480), (1410, 540), label="设备执行", font=label_font)

    img.save(path)


def save_function_diagram(path: Path) -> None:
    img = Image.new("RGB", (1680, 1120), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    box_title_font = load_font(24, bold=True)
    body_font = load_font(20)

    draw.text((60, 36), "功能模块分解图", fill=BLACK, font=title_font)
    add_box(draw, (630, 90, 1060, 220), "OmniHome 产品功能体系", "围绕集中控制、状态感知、自动联动与家庭协同组织业务能力。", title_font=box_title_font, body_font=body_font, fill=ACCENT_LIGHT)

    modules = [
        ((90, 360, 430, 550), "首页总览模块", "展示在线设备、灯光状态、环境指标、异常提示与快捷入口。"),
        ((470, 360, 810, 550), "设备控制模块", "门锁、灯光、空调等设备的开关、调节、执行反馈。"),
        ((850, 360, 1190, 550), "场景与自动化模块", "预设场景执行、启停控制、条件规则配置与联动。"),
        ((1230, 360, 1570, 550), "家庭与通知模块", "家庭成员动态、广播消息、命令历史与提醒回显。"),
        ((280, 760, 620, 950), "门禁与安防模块", "数字钥匙、访客授权、摄像头状态与安防提示。"),
        ((670, 760, 1010, 950), "气候与环境模块", "温湿度、AQI、目标温度、运行模式与环境调节。"),
        ((1060, 760, 1400, 950), "系统支撑模块", "签名校验、同步缓存、异常演示、实时推送与安全控制。"),
    ]
    center = ((630 + 1060) // 2, 220)
    for xy, title, body in modules:
        add_box(draw, xy, title, body, title_font=box_title_font, body_font=body_font)
        target = ((xy[0] + xy[2]) // 2, xy[1])
        draw_arrow(draw, center, target)

    img.save(path)


def save_dfd_diagram(path: Path) -> None:
    img = Image.new("RGB", (1760, 860), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    process_font = load_font(24, bold=True)
    body_font = load_font(19)
    label_font = load_font(18)

    draw.text((60, 36), "DFD 0 层数据流图", fill=BLACK, font=title_font)

    add_small_box(draw, (70, 290, 300, 410), "外部实体\n家庭用户", font=process_font, fill=GREEN_LIGHT)
    add_small_box(draw, (1460, 290, 1690, 410), "外部实体\n智能设备 / 模拟器", font=process_font, fill=GREEN_LIGHT)
    add_box(draw, (640, 210, 1120, 490), "P0 智能家居控制系统", "接收用户控制意图，执行命令编排，维护状态数据，并向客户端回推同步结果。", title_font=process_font, body_font=body_font, fill=ACCENT_LIGHT)
    add_database_symbol(draw, (420, 590, 760, 810), "D1 业务数据库", "设备、房间、场景、自动化、历史记录与全局版本。", title_font=process_font, body_font=body_font)
    add_database_symbol(draw, (980, 590, 1320, 810), "D2 设备状态源", "模拟设备状态、环境读数及可替换的真实设备适配接口。", title_font=process_font, body_font=body_font)

    draw_arrow(draw, (300, 330), (640, 290), label="控制请求 / 查询请求", font=label_font)
    draw_arrow(draw, (640, 420), (300, 380), label="状态总览 / 执行结果 / 通知", font=label_font)
    draw_arrow(draw, (1120, 300), (1460, 330), label="控制命令", font=label_font)
    draw_arrow(draw, (1460, 380), (1120, 430), label="状态反馈 / 事件", font=label_font)
    draw_arrow(draw, (760, 590), (800, 490), label="历史 / 版本 / 主数据", font=label_font)
    draw_arrow(draw, (960, 490), (980, 590), label="设备快照 / 环境数据", font=label_font)

    img.save(path)


def save_sequence_diagram(path: Path) -> None:
    img = Image.new("RGB", (1760, 980), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    lane_font = load_font(22, bold=True)
    text_font = load_font(18)

    draw.text((60, 36), "核心控制链路时序图", fill=BLACK, font=title_font)
    lanes = [
        ("用户", 170),
        ("ArkTS 页面", 480),
        ("控制中心服务", 860),
        ("设备适配层", 1230),
        ("数据库 / WebSocket", 1560),
    ]
    for name, x in lanes:
        add_small_box(draw, (x - 90, 100, x + 90, 170), name, font=lane_font, fill=ACCENT_LIGHT if name != "设备适配层" else ORANGE_LIGHT)
        draw_dashed_line(draw, (x, 180), (x, 910))

    steps = [
        (220, 170, 480, "1. 点击控制设备"),
        (290, 480, 860, "2. 请求签名 / 组装命令"),
        (360, 860, 1230, "3. 校验并调用设备执行"),
        (430, 1230, 860, "4. 返回执行结果"),
        (500, 860, 1560, "5. 写入状态与历史"),
        (570, 1560, 860, "6. 返回当前版本"),
        (640, 860, 1560, "7. 广播 DeviceStateUpdated"),
        (710, 1560, 480, "8. WebSocket 推送"),
        (780, 480, 170, "9. 刷新界面并回显结果"),
    ]
    for y, sx, ex, label in steps:
        draw_arrow(draw, (sx, y), (ex, y), label=label, font=text_font, width=2)

    img.save(path)


def save_er_diagram(path: Path) -> None:
    img = Image.new("RGB", (1760, 980), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    head_font = load_font(22, bold=True)
    body_font = load_font(18)
    label_font = load_font(18)

    draw.text((60, 36), "数据库 ER 简图", fill=BLACK, font=title_font)
    entities = [
        ((90, 150, 440, 360), "rooms", "PK id\nname\nicon\nbuilt_in\nversion\nis_deleted"),
        ((520, 120, 900, 420), "devices", "PK id\nname\nkind\nroom_id(FK)\nstate_json\nupdated_at\nversion\nis_deleted"),
        ((1000, 130, 1370, 380), "scenes", "PK id\nname\ndescription\nenabled\ntrigger_json\ncommands_json\nupdated_at"),
        ((1410, 130, 1670, 380), "automations", "PK id\nname\ntrigger_type\ntrigger_json\naction_json\nenabled\nversion"),
        ((500, 570, 880, 850), "history", "PK id\nrequest_id\ndevice_id(FK)\ncommand_name\nstatus\nmessage\ncreated_at"),
        ((980, 590, 1310, 820), "metadata", "PK key\nvalue\n用于 global_version\nschema_version"),
    ]
    for xy, title, body in entities:
        add_box(draw, xy, title, body.replace("\n", "；"), title_font=head_font, body_font=body_font, fill=LIGHT_GRAY)

    draw_arrow(draw, (440, 250), (520, 250), label="1..n", font=label_font)
    draw_arrow(draw, (900, 260), (1000, 260), label="参与场景", font=label_font)
    draw_arrow(draw, (1370, 250), (1410, 250), label="规则扩展", font=label_font)
    draw_arrow(draw, (710, 420), (710, 570), label="执行留痕", font=label_font)
    draw_arrow(draw, (1140, 590), (1140, 380), label="版本控制", font=label_font)

    img.save(path)


def save_deployment_diagram(path: Path) -> None:
    img = Image.new("RGB", (1720, 940), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(30, bold=True)
    box_title_font = load_font(24, bold=True)
    body_font = load_font(20)
    label_font = load_font(18)

    draw.text((60, 36), "部署拓扑图", fill=BLACK, font=title_font)
    draw.rounded_rectangle((40, 90, 1660, 860), radius=24, outline=BLACK, width=3)
    draw.text((72, 104), "开发 / 演示环境", fill=BLACK, font=box_title_font)

    add_box(draw, (110, 220, 470, 480), "OpenHarmony 设备或模拟器", "运行 ArkTS 客户端，通过局域网地址或 10.0.2.2 访问控制中心接口。", title_font=box_title_font, body_font=body_font, fill=ACCENT_LIGHT)
    add_box(draw, (620, 180, 1040, 520), "控制中心服务主机", "Node.js + Fastify 服务进程，提供 REST API、WebSocket、命令执行、场景与自动化管理。", title_font=box_title_font, body_font=body_font, fill=GREEN_LIGHT)
    add_database_symbol(draw, (1150, 220, 1430, 480), "SQLite", "本地文件 smarthome.db，保存实体主数据、历史记录与版本元数据。", title_font=box_title_font, body_font=body_font)
    add_box(draw, (1150, 560, 1540, 760), "模拟设备 / IoT 适配器", "承担设备执行与状态映射；后续可替换为真实网关、驱动或云侧适配服务。", title_font=box_title_font, body_font=body_font, fill=ORANGE_LIGHT)
    add_box(draw, (620, 620, 1010, 780), "开发工具链", "DevEco Studio、hvigor、npm 与日志调试工具，用于构建、部署与问题定位。", title_font=box_title_font, body_font=body_font)

    draw_arrow(draw, (470, 320), (620, 320), label="HTTP / WebSocket", font=label_font)
    draw_arrow(draw, (1040, 320), (1150, 320), label="SQL 读写", font=label_font)
    draw_arrow(draw, (1040, 430), (1150, 660), label="设备状态映射", font=label_font)
    draw_arrow(draw, (820, 620), (820, 520), label="构建 / 调试", font=label_font)

    img.save(path)


def create_figures() -> dict[str, FigureInfo]:
    ensure_dir(ASSETS_DIR)
    figures = {
        "context": FigureInfo("context", "系统上下文图", ASSETS_DIR / "context.png"),
        "architecture": FigureInfo("architecture", "系统总体架构图", ASSETS_DIR / "architecture.png"),
        "functions": FigureInfo("functions", "功能模块分解图", ASSETS_DIR / "function_modules.png"),
        "dfd": FigureInfo("dfd", "DFD 0 层数据流图", ASSETS_DIR / "dataflow.png"),
        "sequence": FigureInfo("sequence", "核心控制链路时序图", ASSETS_DIR / "sequence.png"),
        "er": FigureInfo("er", "数据库 ER 简图", ASSETS_DIR / "er.png"),
        "deployment": FigureInfo("deployment", "部署拓扑图", ASSETS_DIR / "deployment.png"),
    }
    save_context_diagram(figures["context"].path)
    save_architecture_diagram(figures["architecture"].path)
    save_function_diagram(figures["functions"].path)
    save_dfd_diagram(figures["dfd"].path)
    save_sequence_diagram(figures["sequence"].path)
    save_er_diagram(figures["er"].path)
    save_deployment_diagram(figures["deployment"].path)
    return figures


def configure_run(run, font_name: str, size_pt: float, *, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(paragraph, *, first_line_chars: int = 2, spacing: float = 1.5) -> None:
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(21 * first_line_chars)
    fmt.line_spacing = spacing
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p)
    run = p.add_run(text)
    configure_run(run, "宋体", 12)


def add_bullets(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.left_indent = Pt(24)
        fmt.line_spacing = 1.5
        fmt.space_after = Pt(0)
        run = p.add_run(f"{index}. {item}")
        configure_run(run, "宋体", 12)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(12 if level == 1 else 6)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.5
    size = 16 if level == 1 else 15 if level == 2 else 14
    run = p.add_run(text)
    configure_run(run, "黑体", size)


def add_center_paragraph(doc: Document, text: str, *, font_name: str = "宋体", size: float = 12, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    configure_run(run, font_name, size, bold=bold)


def add_figure_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    configure_run(run, "宋体", 11.5)


def set_table_border(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_cell_text(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    configure_run(run, "宋体", 11 if not bold else 11.5, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, rows: list[list[str]], widths_cm: Iterable[float]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_border(table)
    widths = [Cm(value) for value in widths_cm]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_text(
                cell,
                value,
                bold=row_idx == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            )
    doc.add_paragraph()


def add_figure(doc: Document, figure: FigureInfo, caption: str, *, width_inches: float = 6.35) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(figure.path), width=Inches(width_inches))
    add_figure_caption(doc, caption)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.paragraph_format.space_after = Pt(36)
    run = p.add_run("智能家居控制系统\n产品总体设计文档")
    configure_run(run, "黑体", 22)

    for line in [
        "产品名称：OmniHome 智能家居控制系统",
        "文档类型：正式产品总体设计文档",
        "版本号：V2.0",
        "编制日期：2026年6月26日",
    ]:
        add_center_paragraph(doc, line, size=14)

    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_heading(doc, "目录", 1)
    for line in [
        "1 引言",
        "2 系统总体设计",
        "3 系统架构设计",
        "4 功能模块设计",
        "5 数据设计",
        "6 接口设计",
        "7 运行设计",
        "8 安全设计",
        "9 部署设计",
        "10 设计约束与演进建议",
        "11 总结",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line)
        configure_run(run, "宋体", 12)
    doc.add_page_break()


def build_document() -> None:
    figures = create_figures()
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 编写目的", 2)
    add_body_paragraph(doc, "本文档用于系统化说明 OmniHome 智能家居控制系统的总体设计方案，重点描述系统定位、设计原则、总体架构、功能模块、数据组织、接口体系、运行机制、安全控制与部署方式，为产品评审、课程答辩、后续开发与维护提供统一依据。")
    add_heading(doc, "1.2 项目背景", 2)
    add_body_paragraph(doc, "随着门锁、灯光、空调、摄像头和环境传感器等智能设备不断进入家庭场景，分散式控制入口、设备状态割裂、控制反馈滞后和扩展成本偏高等问题逐步显现。OmniHome 旨在构建统一的家庭控制中枢，通过集中化交互入口和结构化业务支撑，实现设备状态感知、控制执行、场景联动与家庭协同的一体化设计。")
    add_heading(doc, "1.3 建设目标", 2)
    add_bullets(
        doc,
        [
            "构建统一的智能家居控制入口，降低多应用、多页面切换成本。",
            "形成从用户操作、命令执行到状态回写的完整控制闭环。",
            "支持场景与自动化能力，提升系统联动表达能力。",
            "通过共享契约、版本同步和可替换适配层保留工程化扩展空间。",
        ],
    )
    add_heading(doc, "1.4 适用范围", 2)
    add_body_paragraph(doc, "本文档面向产品设计评审、系统实现、测试验证和项目展示场景编制，描述对象为当前仓库中已实现的 OpenHarmony 客户端、Fastify 控制中心服务、SQLite 数据层和模拟设备适配层。文档不将后续可扩展能力表述为当前已落地能力。")

    add_heading(doc, "2 系统总体设计", 1)
    add_heading(doc, "2.1 系统定位", 2)
    add_body_paragraph(doc, "OmniHome 定位为面向家庭场景的集中控制型智能家居原型系统。系统强调以统一控制中枢整合门禁、照明、环境调节、摄像头、场景与自动化等能力，并通过状态总览、命令留痕和实时推送机制支撑可演示、可验证、可演进的产品形态。")
    add_heading(doc, "2.2 设计原则", 2)
    add_table(
        doc,
        [
            ["设计原则", "说明"],
            ["分层解耦", "将表现层、业务层、契约层、数据与设备层分离，降低耦合度。"],
            ["统一契约", "设备模型、命令结构、场景定义与安全信封保持单一语义来源。"],
            ["状态闭环", "控制链路必须覆盖执行结果、状态回写、历史记录与事件广播。"],
            ["可替换扩展", "模拟设备层作为可替换结构，为后续接入真实 IoT 设备保留接口。"],
        ],
        [4.0, 12.5],
    )
    add_figure(doc, figures["context"], "图2-1 系统上下文图")
    add_figure(doc, figures["architecture"], "图2-2 系统总体架构图")
    add_heading(doc, "2.3 系统边界", 2)
    add_body_paragraph(doc, "当前系统边界包括 OpenHarmony 客户端、控制中心服务、共享契约包、SQLite 数据库、设备模拟层以及异常演示接口。商业级账户体系、多租户空间隔离、真实第三方物联网平台对接和云边协同治理能力属于后续演进方向，不纳入本版设计边界。")

    add_heading(doc, "3 系统架构设计", 1)
    add_heading(doc, "3.1 架构概述", 2)
    add_body_paragraph(doc, "系统整体采用“表现层 + 应用服务层 + 共享契约层 + 数据与设备层”的四层架构。表现层负责页面展示与交互触发，应用服务层负责业务编排与规则执行，共享契约层负责统一前后端数据语义，数据与设备层负责状态来源、持久化和执行反馈。")
    add_heading(doc, "3.2 表现层设计", 2)
    add_body_paragraph(doc, "OpenHarmony 客户端基于 ArkTS 页面组织业务入口，通过 ViewModel、Controller 与 Repository 抽离页面渲染逻辑、状态转换逻辑和远程数据访问逻辑，避免界面代码与网络交互和设备语义直接耦合。")
    add_heading(doc, "3.3 应用服务层设计", 2)
    add_body_paragraph(doc, "控制中心服务基于 Fastify 构建，统一装配 devices、access、camera、climate、family、scenes、automations、commands、sync 和 websocket 等业务路由。该层承担命令验签、设备注册表访问、场景执行、历史记录写入、版本递增和事件广播等核心职责。")
    add_heading(doc, "3.4 共享契约层设计", 2)
    add_body_paragraph(doc, "共享契约包负责统一定义设备描述、增强快照、命令结构、场景定义、命令历史和安全信封等关键模型。通过契约前置，系统可以在新增设备类型或扩展控制命令时，以“先扩展模型、再扩展实现”的方式保持演化秩序。")
    add_heading(doc, "3.5 数据与设备层设计", 2)
    add_body_paragraph(doc, "数据与设备层由 SQLite、版本元数据、设备状态快照和模拟设备适配器组成。SQLite 用于支撑设备主数据、场景配置、自动化规则和命令历史的持久化，模拟设备层用于形成可验证的端到端控制闭环，并为未来替换为真实 IoT 网关保留结构位置。")

    add_heading(doc, "4 功能模块设计", 1)
    add_body_paragraph(doc, "系统功能围绕“集中控制、状态感知、联动执行、家庭协同、系统支撑”五类能力展开，模块划分遵循用户认知与业务领域，而非简单按代码目录拆分。")
    add_figure(doc, figures["functions"], "图4-1 功能模块分解图")
    add_table(
        doc,
        [
            ["模块名称", "主要职责", "关键价值"],
            ["首页总览", "展示在线设备、环境数据、灯光数量与异常摘要", "形成统一家庭状态入口"],
            ["设备控制", "门锁、灯光、空调等设备控制与状态反馈", "完成高频核心控制闭环"],
            ["场景与自动化", "预设场景执行、规则启停与联动管理", "提升家庭智能联动能力"],
            ["家庭与通知", "家庭成员动态、广播消息、命令历史回显", "增强家庭协同与留痕"],
            ["系统支撑", "签名校验、同步接口、WebSocket 推送、异常演示", "保障系统可靠性与可验证性"],
        ],
        [3.2, 8.2, 4.8],
    )
    add_heading(doc, "4.1 首页总览模块", 2)
    add_body_paragraph(doc, "首页总览模块作为系统入口页面，重点承载设备在线概览、灯光状态统计、温湿度与 AQI 摘要、风险提示和各领域页面的快捷入口。设计目标不是罗列全部细节，而是让用户在最短时间内把握家庭总体运行状态。")
    add_heading(doc, "4.2 设备控制模块", 2)
    add_body_paragraph(doc, "设备控制模块面向门锁、灯光和空调等核心家居设备，统一采用“读取当前状态、发起控制请求、执行结果回显、历史记录留痕”的交互模式。该模块体现了系统最核心的业务链路，也是正式设计中必须优先保证一致性和可靠性的部分。")
    add_heading(doc, "4.3 场景与自动化模块", 2)
    add_body_paragraph(doc, "场景模块负责显式触发的一组设备动作组合，自动化模块负责基于触发条件执行规则。两者在语义层区分“用户主动执行”和“系统条件触发”，既便于功能表达，也为后续扩展更复杂的规则引擎保留清晰边界。")
    add_heading(doc, "4.4 家庭与通知模块", 2)
    add_body_paragraph(doc, "家庭与通知模块支撑家庭成员动态、广播消息、命令执行记录和提醒回显等能力，使系统不局限于单一设备控制，而具备家庭协同视角。该模块对于产品演示和后续产品化延展具有较高价值。")

    add_heading(doc, "5 数据设计", 1)
    add_heading(doc, "5.1 数据设计目标", 2)
    add_body_paragraph(doc, "数据设计既服务于系统运行时状态管理，也服务于命令留痕、版本同步和可扩展性控制。当前方案采用轻量级 SQLite 作为本地持久化介质，以满足原型阶段的低部署成本和高可验证性要求。")
    add_figure(doc, figures["er"], "图5-1 数据库 ER 简图")
    add_heading(doc, "5.2 主要数据实体", 2)
    add_table(
        doc,
        [
            ["数据实体", "核心字段", "设计说明"],
            ["metadata", "key、value", "保存 global_version、schema_version 等元数据。"],
            ["rooms", "id、name、icon、version", "定义房间信息，支持设备归属与页面分组。"],
            ["devices", "id、kind、room_id、state_json、version", "记录设备快照，是总览和控制页面的数据基础。"],
            ["scenes", "id、name、enabled、commands_json", "保存场景定义和场景动作集合。"],
            ["automations", "id、trigger_json、action_json、enabled", "保存自动化规则及触发动作表达。"],
            ["history", "request_id、device_id、command_name、status", "记录命令执行结果，支撑通知与审计。"],
        ],
        [3.0, 6.2, 7.0],
    )
    add_heading(doc, "5.3 版本化设计", 2)
    add_body_paragraph(doc, "系统通过 metadata.global_version 和各实体 version 字段实现增量同步。任一关键业务对象变更后，服务端在事务中同步提升版本号，使客户端可以只拉取发生变化的数据集合，降低全量刷新带来的资源开销和状态不一致风险。")

    add_heading(doc, "6 接口设计", 1)
    add_heading(doc, "6.1 接口分层", 2)
    add_body_paragraph(doc, "系统接口按照业务领域进行划分，分别覆盖设备总览、命令控制、场景与自动化、垂直领域服务、同步服务和实时推送服务。这种分层方式便于客户端按模块装配能力，也便于服务端按职责演进。")
    add_table(
        doc,
        [
            ["接口分类", "代表接口", "说明"],
            ["设备与摘要接口", "GET /api/devices；GET /api/devices/:id；GET /api/summary", "提供首页、设备详情和概览页面所需状态数据。"],
            ["控制接口", "POST /api/demo/sign-command；POST /api/commands；GET /api/commands/history", "实现命令签名、命令执行和结果追踪。"],
            ["场景与自动化接口", "GET/PATCH/POST /api/scenes；GET/POST/PUT/DELETE /api/automations", "支持规则查询、启停和执行。"],
            ["垂直领域接口", "GET /api/access；GET /api/cameras；GET /api/family；GET /api/climate", "服务门禁、摄像头、家庭和气候模块。"],
            ["同步与实时接口", "GET /api/sync?lastVersion=n；GET /ws/events", "实现按版本补偿同步与状态实时推送。"],
        ],
        [3.4, 7.4, 5.4],
    )
    add_heading(doc, "6.2 接口设计要点", 2)
    add_bullets(
        doc,
        [
            "控制类接口采用先签名后执行的双阶段模式，避免客户端直接组装不受保护的命令报文。",
            "同步接口以版本号为中心组织差量返回结果，提升状态同步效率。",
            "WebSocket 作为主动推送通道，用于补充页面轮询并缩短用户反馈时延。",
        ],
    )

    add_heading(doc, "7 运行设计", 1)
    add_body_paragraph(doc, "运行设计重点描述从用户意图触发到界面状态回显的全链路过程，覆盖控制链路、数据链路、同步链路和实时链路。")
    add_figure(doc, figures["dfd"], "图7-1 DFD 0 层数据流图")
    add_figure(doc, figures["sequence"], "图7-2 核心控制链路时序图")
    add_heading(doc, "7.1 读取链路", 2)
    add_body_paragraph(doc, "客户端进入首页或领域页面后，通过 Repository 发起摘要查询或领域查询请求。服务端从数据库和设备状态源装配结果并返回前端，前端再将返回结构映射为页面展示模型。")
    add_heading(doc, "7.2 控制链路", 2)
    add_body_paragraph(doc, "用户发起设备控制后，前端首先请求签名接口构造安全信封，再调用命令执行接口提交控制意图。服务端完成验签、重放保护、设备查找和执行后，更新数据库状态和历史记录，并将结果反馈客户端。")
    add_heading(doc, "7.3 同步链路", 2)
    add_body_paragraph(doc, "客户端保存本地 lastVersion，同步接口根据该版本号返回增量变更的设备、房间、场景与自动化数据。该模式兼顾性能和一致性，是系统支撑多页面刷新与后台状态变化的关键机制。")
    add_heading(doc, "7.4 实时链路", 2)
    add_body_paragraph(doc, "当设备状态变化、场景执行完成或异常演示触发时，服务端通过 WebSocket 广播事件。客户端接收事件后对本地缓存和页面模型进行局部刷新，从而实现接近实时的状态回显。")

    add_heading(doc, "8 安全设计", 1)
    add_heading(doc, "8.1 命令签名机制", 2)
    add_body_paragraph(doc, "系统通过 HMAC-SHA256 对控制命令进行签名，安全信封包含 requestId、timestamp、nonce、payload 摘要等字段。客户端不直接伪造执行报文，而是以服务端下发的受保护结构作为命令提交载体。")
    add_heading(doc, "8.2 重放保护机制", 2)
    add_body_paragraph(doc, "服务端借助 ReplayGuard 记录 nonce 和时间窗口，对重复命令和超时命令进行拒绝，降低控制链路受到重放攻击的风险。该机制对于门锁类高敏感设备控制尤为关键。")
    add_heading(doc, "8.3 异常演示与安全反馈", 2)
    add_body_paragraph(doc, "系统提供安全拒绝演示接口，用于在原型环境中展示授权失败、设备离线等异常状态下的用户可见反馈。正式设计中将异常反馈纳入主链路，有助于提升系统的可解释性和可验证性。")
    add_heading(doc, "8.4 传输安全预留", 2)
    add_body_paragraph(doc, "控制中心服务预留 TLS 证书与私钥配置项，支持后续切换至 HTTPS 监听模式。当前版本的重点在于将命令签名、重放保护和异常反馈等机制纳入体系化设计，而非停留在后期补丁式加固。")

    add_heading(doc, "9 部署设计", 1)
    add_body_paragraph(doc, "系统采用轻量级开发与演示部署结构，由 OpenHarmony 客户端、Node.js 控制中心服务、SQLite 数据库以及模拟设备 / 适配层组成，部署成本低、链路清晰，适合课程展示与功能验证。")
    add_figure(doc, figures["deployment"], "图9-1 部署拓扑图")
    add_table(
        doc,
        [
            ["部署单元", "运行环境", "说明"],
            ["OpenHarmony 客户端", "模拟器或真机", "运行 ArkTS 页面逻辑并访问控制中心服务。"],
            ["控制中心服务", "Node.js 主机进程", "提供 REST API、WebSocket、签名与业务编排能力。"],
            ["SQLite 数据库", "服务主机本地文件", "保存业务主数据、历史记录与版本元数据。"],
            ["模拟设备 / 适配层", "服务进程内或适配模块", "用于演示设备执行与后续真实设备接入。"],
        ],
        [3.5, 4.6, 8.1],
    )
    add_heading(doc, "9.1 客户端部署说明", 2)
    add_body_paragraph(doc, "在模拟器环境下，客户端通常通过 10.0.2.2 访问宿主机上的控制中心服务；在局域网真机环境下，可替换为实际服务主机地址。该设计兼容演示和开发调试两类典型场景。")
    add_heading(doc, "9.2 服务端部署说明", 2)
    add_body_paragraph(doc, "控制中心服务依赖 Node.js、Fastify、better-sqlite3 和共享契约包运行。服务启动后同时提供 REST API 和 WebSocket 通道，以支持页面读取、命令控制、同步补偿和实时推送等能力。")

    add_heading(doc, "10 设计约束与演进建议", 1)
    add_heading(doc, "10.1 当前设计约束", 2)
    add_bullets(
        doc,
        [
            "当前设备层主要依赖模拟器与适配抽象，尚未接入真实商业 IoT 平台。",
            "SQLite 适合原型与中低并发场景，大规模并发下需要进一步演进为独立数据库服务。",
            "身份认证、权限体系和多家庭空间隔离能力尚未纳入本版正式边界。",
        ],
    )
    add_heading(doc, "10.2 演进建议", 2)
    add_bullets(
        doc,
        [
            "引入真实设备网关或驱动适配层，保持上层控制接口不变，逐步替换模拟设备实现。",
            "扩展规则引擎、用户身份认证与权限控制，提升产品化落地能力。",
            "将同步、事件与审计能力进一步服务化，支撑多终端和更高复杂度场景。",
        ],
    )

    add_heading(doc, "11 总结", 1)
    add_body_paragraph(doc, "OmniHome 智能家居控制系统在当前版本中已经形成较为完整的正式产品总体设计：其一，系统边界、核心职责和分层关系明确；其二，控制闭环、同步机制和安全机制形成统一主线；其三，图示化表达覆盖了上下文、架构、功能、数据流、时序、数据实体和部署拓扑等关键视角。文档可直接作为产品设计评审、课程答辩和后续扩展实现的基础资料。")

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    ensure_dir(ASSETS_DIR)
    build_document()
